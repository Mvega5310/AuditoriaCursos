"""Genera la plantilla de Excel que descarga el cliente (con instrucciones y ejemplos de fechas relativas a hoy)."""
import io
from datetime import timedelta

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import hoy

_ENCABEZADO = PatternFill("solid", fgColor="0B5CAD")

INSTRUCCIONES = [
    ("Cómo preparar su archivo", True),
    ("", False),
    ("Solo necesita UNA de estas hojas (puede usar las dos). El nombre de la hoja no importa: se reconocen por sus columnas.", False),
    ("", False),
    ("1) LISTA  (hoja 'Documentos'): una fila por documento. Obligatorio: la columna 'Documento' (o 'Curso') y la 'Fecha Vencimiento'.", False),
    ("    Sin cédula = documento de la EMPRESA (p. ej. una póliza). Con cédula = documento de una PERSONA.", False),
    ("2) MATRIZ  (hoja 'Matriz'): una fila por persona y una columna por curso, con la fecha en cada celda.", False),
    ("    Obligatorio: 'Cedula' y 'Nombre'. El encabezado del curso puede llevar su vigencia en días: 'RCP | 730'.", False),
    ("", False),
    ("Lo demás NO estorba:", True),
    ("  - Columnas extra (teléfono, observaciones, consecutivos…) se ignoran.", False),
    ("  - Los títulos o filas vacías encima de la tabla se saltan solos.", False),
    ("  - Una celda vacía, 'N/A', '-' o 'Pendiente' significa 'sin registrar'.", False),
    ("  - Una fila con un dato malo se omite y se le informa cuál era; las demás sí se guardan.", False),
    ("", False),
    ("Fechas: dd/mm/aaaa (por ejemplo 15/03/2027) o aaaa-mm-dd.", False),
    ("Antes de guardar puede usar el botón 'Validar sin guardar' para ver el resultado sin cambiar nada.", False),
    ("Puede subir el archivo varias veces: actualiza sin duplicar.", False),
    ("", False),
    ("Exámenes médicos: registre solo fechas y vigencia. No incluya diagnósticos ni resultados (Ley 1581 de 2012).", False),
]


def _hoja(wb, nombre, encabezados, filas, anchos=None):
    ws = wb.create_sheet(nombre)
    ws.append(encabezados)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _ENCABEZADO
    for f in filas:
        ws.append(f)
    for i, enc in enumerate(encabezados, 1):
        ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(i, max(14, len(str(enc)) + 4))
    ws.freeze_panes = "A2"
    return ws


def construir_plantilla() -> bytes:
    h = hoy()
    f = lambda d: (h + timedelta(days=d)).strftime("%d/%m/%Y")      # noqa: E731
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    for texto, negrita in INSTRUCCIONES:
        ws.append([texto])
        ws.cell(ws.max_row, 1).font = Font(bold=negrita, size=13 if negrita else 11)
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110

    _hoja(wb, "Documentos",
          ["Cedula", "Nombre", "Apellido", "Cargo", "Area", "Documento", "Referencia", "Fecha Emision", "Fecha Vencimiento"],
          [["1001", "Ana", "Pérez", "Auxiliar de enfermería", "Urgencias", "RCP", "", f(-300), f(65)],
           ["1002", "Luis", "Gómez", "Soldador", "Planta", "Trabajo en alturas", "CERT-88", "", f(20)],
           ["", "", "", "", "", "Póliza todo riesgo", "POL-123", "", f(200)]])

    _hoja(wb, "Matriz",
          ["Cedula", "Nombre", "Cargo", "Area", "RCP | 730", "Bioseguridad | 365", "Observaciones"],
          [["2001", "María Rojas", "Enfermera", "UCI", f(-700), f(-360), "se ignora esta columna"],
           ["2002", "Pedro Díaz", "Enfermero", "Urgencias", f(-740), "", ""]])

    _hoja(wb, "Requisitos", ["Cargo", "Documento"],
          [["Enfermera", "RCP"], ["Enfermera", "Bioseguridad"], ["*", "Afiliación ARL"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
