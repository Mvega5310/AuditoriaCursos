"""Formato de hospital: encabezado de varias filas con celdas combinadas, pares inicio/final,
cursos controlados, N/A = no aplica, hoja de retirados. Datos sinteticos (nunca reales)."""
import sys
from datetime import timedelta
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "scripts"))

import core                      # noqa: E402
import importar_excel            # noqa: E402
import generar_alertas           # noqa: E402

H = core.hoy()


def _f(dias):
    return (H + timedelta(days=dias)).strftime("%d/%m/%Y")


def _libro(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AUXILIARES"
    ws.append(["", "", "", "", "BLS", "", "GESTION DEL DUELO", "DENGUE", "UCI (2 AÑOS)", ""])
    ws.append(["CEDULA", "NOMBRE", "CARGO", "UBICACION", "FECHA INICIAL", "FECHA FINAL",
               "FECHA ACTUALIZACION", "FECHA ACTUALIZACION", "FECHA INICIAL", "FECHA FINAL"])
    ws.merge_cells("E1:F1")
    ws.merge_cells("I1:J1")
    ws.append(["10010001", "Ana Perez", "Auxiliar de enfermeria", "UCI", "N/A", "N/A", _f(100), _f(-5), _f(-800), ""])
    ws.append(["10020002", "Luis Ruiz", "Auxiliar de enfermeria", "Urgencias", _f(-30), _f(700), "X:" + _f(-3), _f(-5), "N/A", "N/A"])
    ws.append(["10030003", "Eva Gil", "Auxiliar de enfermeria", "Urgencias", "", "", "", "", "", ""])
    ws.append(["10040004", "Raul Mora", "Auxiliar de enfermeria", "Urgencias", "", "", "", "", "", ""])
    ret = wb.create_sheet("RETIRADOS")
    ret.append(["10030003", "Eva Gil"])
    ruta = tmp_path / "huc.xlsx"
    wb.save(ruta)
    return ruta


def _importar(tmp_path, **kw):
    db = tmp_path / "t.db"
    core.inicializar(db) if hasattr(core, "inicializar") else None
    stats = importar_excel.importar(_libro(tmp_path), db_path=db, verbose=False, empresa_forzada="Hospital",
                                    fechas="vencimiento", solo_cursos=["BLS", "Gestion del duelo", "UCI"],
                                    exigir_cursos=True, **kw)
    return db, stats


def test_lee_pares_na_y_solo_controlados(tmp_path):
    db, stats = _importar(tmp_path)
    conn = core.conectar(db)
    docs = {(r["cedula"], r["nombre"]): r for r in conn.execute("""
        SELECT e.cedula, t.nombre, d.fecha_emision, d.fecha_vencimiento, d.no_aplica
        FROM documentos d JOIN empleados e ON e.id = d.empleado_id JOIN tipos_documento t ON t.id = d.tipo_id""")}
    nombres = {n for _, n in docs}
    assert not any("engue" in n for n in nombres)                       # no controlado
    ana_bls = next(v for (c, n), v in docs.items() if c == "10010001" and "BLS" in n.upper())
    assert ana_bls["no_aplica"] == 1
    luis_duelo = next(v for (c, n), v in docs.items() if c == "10020002" and "duelo" in n.lower())
    assert luis_duelo["fecha_vencimiento"] == (H - timedelta(days=3)).isoformat()   # "X:fecha"
    assert not any(c == "10010001" and n.upper().startswith("UCI") for (c, n) in docs)   # inicio sin final: descartado
    assert stats["sin_fecha_final"] == {"UCI": 1}
    assert stats["no_aplica"] >= 2


def test_faltantes_na_no_cuenta_y_vacio_si(tmp_path):
    db, _ = _importar(tmp_path)
    conn = core.conectar(db)
    eid = conn.execute("SELECT id FROM empresas WHERE nombre = 'Hospital'").fetchone()["id"]
    falt = {f["identificacion"]: f for f in generar_alertas.consultar_faltantes(conn, eid)}
    assert "BLS" not in falt["10010001"]["faltan"].upper()                     # N/A no es faltante
    assert "UCI" in falt["10010001"]["faltan"].upper()                         # su UCI se descarto
    assert "10020002" not in falt                                              # Luis: todo registrado o N/A
    assert "10030003" not in falt                                              # Eva esta retirada: inactiva
    assert "BLS" in falt["10040004"]["faltan"].upper()                         # celda en blanco = pendiente


def test_retirados_quedan_inactivos(tmp_path):
    db, stats = _importar(tmp_path)
    conn = core.conectar(db)
    eva = conn.execute("SELECT activo FROM empleados WHERE cedula = '10030003'").fetchone()
    assert eva["activo"] == 0
    assert stats["retirados"] == 1


def test_fecha_invalida_se_reporta_con_hoja_fila_y_persona(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ENFERMERAS"
    ws.append(["CEDULA", "NOMBRE", "CARGO", "BLS"])
    ws.append(["10050005", "Marta Lopez", "Enfermera", "31/02/2025"])
    ws.append(["10060006", "Juan Soto", "Enfermero", _f(200)])
    ruta = tmp_path / "x.xlsx"
    wb.save(ruta)
    stats = importar_excel.importar(ruta, db_path=tmp_path / "e.db", verbose=False, empresa_forzada="Hospital",
                                    fechas="vencimiento")
    msg = " ".join(stats["errores"])
    assert "ENFERMERAS" in msg and "Marta Lopez" in msg and "10050005" in msg and "BLS" in msg and "31/02/2025" in msg
