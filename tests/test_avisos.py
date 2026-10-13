"""Ventanas de alerta (periodo + anticipacion), avisos por empresa, panel con filtros, Excel y resumen al importar."""
import io
import sys
from datetime import timedelta
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "scripts"))

import config                    # noqa: E402
import core                      # noqa: E402
import generar_alertas           # noqa: E402
from test_web import app, registrar, subir, _csrf   # noqa: E402,F401

H = core.hoy()


def _f(d):
    return (H + timedelta(days=d)).strftime("%d/%m/%Y")


def _excel(filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cedula", "Nombre", "Apellido", "Cargo", "Area", "Documento", "Fecha Vencimiento"])
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


FILAS = [
    ["101", "Ana", "Uno", "Enfermera", "UCI", "BLS", _f(-5)],          # vencido
    ["101", "Ana", "Uno", "Enfermera", "UCI", "ACLS", _f(5)],          # 7
    ["102", "Luis", "Dos", "Auxiliar", "Urgencias", "BLS", _f(12)],    # 15
    ["102", "Luis", "Dos", "Auxiliar", "Urgencias", "ACLS", _f(25)],   # 30
    ["103", "Eva", "Tres", "Auxiliar", "Urgencias", "BLS", _f(40)],    # solo semanal 37? no: 40 > 37
    ["104", "Raul", "Cuatro", "Auxiliar", "UCI", "BLS", _f(55)],       # mensual 60
]


def test_ventana_es_periodo_mas_anticipacion():
    assert config.ventana_alerta("diaria", 90) == 7                  # fija
    assert config.ventana_alerta("semanal") == 37
    assert config.ventana_alerta("quincenal") == 45
    assert config.ventana_alerta("mensual") == 60
    assert config.ventana_alerta("quincenal", 60) == 75


def test_resumen_al_importar_y_validar(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    r = c.post("/importar", data={"csrf": _csrf(c, "/importar"), "archivo": (_excel(FILAS), "d.xlsx"),
                                  "accion": "validar"}, content_type="multipart/form-data")
    html = r.data.decode()
    assert "Así quedaría Hospital Z" in html
    assert ">1</b>Vencidos" in html and ">1</b>Vencen en 7" in html and ">2</b>Vencen en 15" in html
    assert ">3</b>Vencen en 30" in html


def test_panel_filtra_por_ventana_curso_y_ubicacion(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    subir(c, _excel(FILAS))
    assert "6 registros" in c.get("/panel").data.decode()                       # por defecto: ventana maxima (60)
    assert "1 registro" in c.get("/panel?ventana=vencidos").data.decode()
    assert "3 registros" in c.get("/panel?ventana=15").data.decode()
    html = c.get("/panel?ventana=60&curso=BLS&ubicacion=UCI").data.decode()
    assert "2 registros" in html and "Raul" in html and "Luis" not in html.split("Documentos exigidos")[0]
    assert c.get("/panel?ventana=abc&curso=<script>").status_code == 200        # valores raros: se ignoran
    x = c.get("/panel.xlsx?ventana=30&ubicacion=Urgencias")
    ws = openpyxl.load_workbook(io.BytesIO(x.data))["Vencimientos"]
    assert [r[0].value for r in ws.iter_rows(min_row=2)] == ["Dos, Luis", "Dos, Luis"]


def test_avisos_por_empresa(app, tmp_path, monkeypatch):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    subir(c, _excel(FILAS))
    r = c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "10", "alertas": ["quincenal"]})
    assert r.status_code == 302
    assert "15 + 10 = <b>25 días" in c.get("/avisos").data.decode()
    assert c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "999"}).status_code == 302
    conn = core.conectar(app.config["DB_PATH"])
    eid = conn.execute("SELECT id FROM empresas WHERE nombre = 'Hospital Z'").fetchone()["id"]
    conf = generar_alertas.configuracion_empresa(conn, eid)
    assert conf["anticipacion"] == 10 and conf["activas"] == {"quincenal"}          # el 999 no se guardo

    salidas = tmp_path / "salidas"
    monkeypatch.setattr(generar_alertas, "SALIDAS_DIR", salidas)
    generar_alertas.ejecutar_alerta("semanal", dry_run=True, db_path=app.config["DB_PATH"])
    assert not salidas.exists() or not list(salidas.glob("*hospital_z*"))            # no recibe semanal
    generar_alertas.ejecutar_alerta("quincenal", dry_run=True, db_path=app.config["DB_PATH"])
    html = (salidas / "alerta_quincenal_hospital_z.html").read_text(encoding="utf-8")
    assert "próximos 25 días" in html and "Cuatro" not in html and "Dos" in html     # 55 dias queda fuera

    c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "30"})   # ninguna marcada
    assert generar_alertas.configuracion_empresa(conn, eid)["activas"] == set()


def test_excel_del_panel_no_ejecuta_formulas(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cedula", "Nombre", "Apellido", "Cargo", "Area", "Documento", "Fecha Vencimiento"])
    ws.append(["105", "=HYPERLINK(\"http://x\",\"clic\")", "", "Auxiliar", "UCI", "BLS", _f(3)])
    ws["B2"].data_type = "s"                       # texto que empieza por "=", como lo guardaria un usuario
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    subir(c, buf)
    x = c.get("/panel.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(x.data))
    celda = wb["Vencimientos"]["A2"]
    assert celda.data_type == "s" and celda.value.startswith("=HYPERLINK")


# ------------------------------------------------------------ dia del aviso semanal

def _conf(app, nombre="Hospital Z"):
    conn = core.conectar(app.config["DB_PATH"])
    eid = conn.execute("SELECT id FROM empresas WHERE nombre = ?", (nombre,)).fetchone()["id"]
    return conn, eid


def test_dia_semanal_por_defecto_y_se_puede_elegir(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    conn, eid = _conf(app)
    assert generar_alertas.configuracion_empresa(conn, eid)["dia_semanal"] == 0            # lunes
    c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "30", "dia_semanal": "2",
                            "alertas": ["semanal"]})
    assert generar_alertas.configuracion_empresa(conn, eid)["dia_semanal"] == 2
    assert "Semanal (miércoles)" in c.get("/avisos").data.decode()


def test_dia_semanal_invalido_no_se_guarda_y_si_falta_se_conserva(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    conn, eid = _conf(app)
    c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "30", "dia_semanal": "3"})
    for malo in ("7", "-1", "sabado"):
        c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "30", "dia_semanal": malo})
        assert generar_alertas.configuracion_empresa(conn, eid)["dia_semanal"] == 3
    c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "20"})            # formulario antiguo, sin el campo
    assert generar_alertas.configuracion_empresa(conn, eid)["dia_semanal"] == 3


def test_programador_envia_el_semanal_solo_en_el_dia_de_cada_empresa(app, tmp_path, monkeypatch):
    for nombre, correo in (("Hospital Z", "z@z.com"), ("Clinica Y", "y@y.com")):
        c, _ = registrar(app, nombre, correo)
        subir(c, _excel(FILAS))
    conn, _ = _conf(app)
    conn.execute("UPDATE empresas SET dia_semanal = 0 WHERE nombre = 'Hospital Z'")        # lunes
    conn.execute("UPDATE empresas SET dia_semanal = 2 WHERE nombre = 'Clinica Y'")         # miercoles
    conn.commit()
    salidas = tmp_path / "salidas"
    monkeypatch.setattr(generar_alertas, "SALIDAS_DIR", salidas)

    def corre(dia_iso, **kw):                      # semana sin festivos: 2026-10-19 es lunes
        monkeypatch.setattr(generar_alertas, "hoy", lambda: __import__("datetime").date(2026, 10, 19 + dia_iso))
        for f in salidas.glob("*") if salidas.exists() else []:
            f.unlink()
        generar_alertas.ejecutar_alerta("semanal", dry_run=True, db_path=app.config["DB_PATH"], **kw)
        return sorted(p.name for p in salidas.glob("*")) if salidas.exists() else []

    assert corre(0, respetar_dia=True) == ["alerta_semanal_hospital_z.html"]               # lunes
    assert corre(1, respetar_dia=True) == []                                               # martes: nadie
    assert corre(2, respetar_dia=True) == ["alerta_semanal_clinica_y.html"]                # miercoles
    assert len(corre(3)) == 2                                                              # manual: todas, sin filtrar por dia


def test_el_programador_revisa_el_semanal_de_lunes_a_viernes():
    import scheduler
    from apscheduler.schedulers.background import BackgroundScheduler
    sch = BackgroundScheduler(timezone="America/Bogota")
    scheduler.registrar_jobs(sch)
    job = sch.get_job("alerta_semanal")
    assert str(job.trigger.fields[job.trigger.FIELD_NAMES.index("day_of_week")]) == "mon-fri"
    assert job.kwargs == {"respetar_dia": True}


def test_avisos_muestra_el_proximo_envio_semanal(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    html = c.get("/avisos").data.decode()
    assert "Próximo aviso semanal:" in html and "festivos de Colombia" in html
    c.post("/avisos", data={"csrf": _csrf(c, "/avisos"), "anticipacion": "30", "dia_semanal": "1", "alertas": ["mensual"]})
    assert "Próximo aviso semanal:" not in c.get("/avisos").data.decode()          # semanal desactivado: no se anuncia


def test_motor_aplica_el_calendario_a_diaria_quincenal_y_mensual(app, tmp_path, monkeypatch):
    from datetime import date
    import calendario
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    subir(c, _excel(FILAS))
    salidas = tmp_path / "salidas"
    monkeypatch.setattr(generar_alertas, "SALIDAS_DIR", salidas)
    original = calendario.toca_hoy

    def corre(tipo, fecha_calendario):
        # solo la FECHA del calendario se simula; los documentos se evaluan con la fecha real
        monkeypatch.setattr(calendario, "toca_hoy", lambda t, _h, dia=0: original(t, fecha_calendario, dia))
        for f in (salidas.glob("*") if salidas.exists() else []):
            f.unlink()
        generar_alertas.ejecutar_alerta(tipo, dry_run=True, db_path=app.config["DB_PATH"], respetar_dia=True)
        return sorted(p.name for p in salidas.glob("*")) if salidas.exists() else []

    assert corre("diaria", date(2026, 10, 12)) == []                                        # festivo
    assert corre("diaria", date(2026, 10, 10)) == []                                        # sabado
    assert corre("diaria", date(2026, 10, 13)) == ["alerta_diaria_hospital_z.html"]
    assert corre("mensual", date(2026, 8, 1)) == []                                         # 1 de agosto: sabado
    assert corre("mensual", date(2026, 8, 3)) == ["alerta_mensual_hospital_z.html"]         # sale el lunes 3
    assert corre("mensual", date(2026, 8, 4)) == []
    assert corre("quincenal", date(2026, 10, 14)) == []
    assert corre("quincenal", date(2026, 10, 15)) == ["alerta_quincenal_hospital_z.html"]
    # una ejecucion manual (sin respetar_dia) no consulta el calendario
    generar_alertas.ejecutar_alerta("mensual", dry_run=True, db_path=app.config["DB_PATH"])
    assert (salidas / "alerta_mensual_hospital_z.html").exists()


# ------------------------------------------------------------ panel: secciones desplegables y paginacion de 20

def _excel_45():
    """45 personas con un BLS que vence en 10 dias, y el requisito ARL para todos (45 faltantes)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(["Cedula", "Nombre", "Apellido", "Cargo", "Area", "Documento", "Fecha Vencimiento"])
    for i in range(1, 46):
        ws.append([str(1000 + i), f"Persona{i:02d}", "Prueba", "Auxiliar", "UCI" if i % 2 else "Urgencias", "BLS", _f(10)])
    req = wb.create_sheet("Requisitos")
    req.append(["Cargo", "Documento"])
    req.append(["*", "Afiliación ARL"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _filas(html, seccion):
    """Cantidad de filas de datos dentro de la seccion desplegable `seccion`."""
    parte = html.split(f'id="{seccion}"', 1)[1].split("</details>", 1)[0]
    return parte.count("<tr><td>")


def test_panel_pagina_de_a_20_cada_lista(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    subir(c, _excel_45())
    h1 = c.get("/panel").data.decode()
    assert "45 registros" in h1 and _filas(h1, "vencimientos") == 20 and _filas(h1, "faltantes") == 20
    assert "Mostrando 1–20 de 45" in h1 and "Siguiente ›" in h1
    h3 = c.get("/panel?pagina=3").data.decode()
    assert _filas(h3, "vencimientos") == 5 and "Mostrando 41–45 de 45" in h3 and "Siguiente ›" not in h3.split('id="faltantes"')[0]
    assert _filas(c.get("/panel?pf=3").data.decode(), "faltantes") == 5
    for raro in ("0", "-1", "999", "abc"):                       # valores raros: pagina valida, sin error
        r = c.get(f"/panel?pagina={raro}")
        assert r.status_code == 200 and _filas(r.data.decode(), "vencimientos") in (20, 5)
    # el Excel trae la lista completa, no solo la pagina
    ws = openpyxl.load_workbook(io.BytesIO(c.get("/panel.xlsx?pagina=2").data))["Vencimientos"]
    assert ws.max_row - 1 == 45


def test_panel_paginacion_conserva_filtros_y_secciones(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    subir(c, _excel_45())
    h = c.get("/panel?ubicacion=UCI&pf=2").data.decode()
    assert "23 registros" in h                                   # 23 impares en UCI: 2 paginas
    assert "ubicacion=UCI" in h.split('id="vencimientos"')[1].split("</details>")[0].split("<nav")[1]
    assert '<details class="card desplegable" id="vencimientos" open>' in h
    assert '<details class="card desplegable" id="faltantes" open>' in h          # se esta paginando: abierta
    h2 = c.get("/panel").data.decode()
    assert '<details class="card desplegable" id="faltantes" >' in h2                # por defecto: plegada


def test_importar_requisitos_del_archivo_plegados(app):
    c, _ = registrar(app, "Hospital Z", "z@z.com")
    h = c.get("/importar").data.decode()
    assert '<details class="card desplegable" id="requisitos-archivo">' in h        # sin "open": plegada
    assert "¿Qué necesita su archivo?" in h and "Columnas obligatorias" in h       # el contenido sigue en la pagina
